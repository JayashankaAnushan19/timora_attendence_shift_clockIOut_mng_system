"""Employee types, pay details, added overtime and who may use them."""
import io
import sqlite3

import pytest
from openpyxl import load_workbook

from backend.database import SCHEMA_VERSION, init_db
from .conftest import Api

RANGE = "preset=custom&from=2026-10-26&to=2026-10-26"


@pytest.fixture
def login_as(app, admin):
    roles = {r["name"]: r["id"] for r in admin.get("/api/admin/roles").json["items"]}

    def create(role, username):
        assert admin.post("/api/admin/users", {
            "username": username, "password": "Temp-pass1", "role_id": roles[role]}).status_code == 201
        api = Api(app.test_client())
        assert api.login(username, "Temp-pass1").status_code == 200
        assert api.post("/api/auth/change-password", {"current_password": "Temp-pass1", "new_password": "Own-pass-99"}).status_code == 200
        return api
    return create


def add_ot(api, codes, hours=2, date="2026-10-26", reason="Urgent repair"):
    return api.post("/api/admin/overtime", {"employee_ids": codes, "work_date": date, "hours": hours, "reason": reason})


# --- Employee types and pay details ---------------------------------------

def test_employee_types_and_filter(make_employee, admin):
    make_employee("Tech", employee_type="Technician")
    make_employee("Boss", employee_type="Manager")
    assert admin.get("/api/admin/employees?employee_type=Manager").json["items"][0]["name"] == "Boss"
    assert admin.post("/api/admin/employees", {"name": "X", "department": "D", "designation": "D", "employee_type": "Astronaut"}).json["field"] == "employee_type"
    assert admin.get("/api/admin/lookups").json["employee_types"] == ["Technician", "Manager", "Cleaner"]


def test_employee_types_setting_can_be_changed(make_employee, admin):
    assert admin.put("/api/admin/settings", {"employee_types": ["Driver", "Cleaner"]}).status_code == 200
    assert make_employee("Dee", employee_type="Driver")
    assert admin.put("/api/admin/settings", {"employee_types": []}).status_code == 400
    assert admin.put("/api/admin/settings", {"employee_types": ["A", "a"]}).status_code == 400


def test_salary_and_rate_saved_exactly(make_employee, admin):
    code = make_employee("Paid", basic_salary="85000.50", ot_rate=650)
    employee = admin.get(f"/api/admin/employees/{code}").json["employee"]
    assert employee["basic_salary"] == 85000.5 and employee["ot_rate"] == 650
    assert "basic_salary_cents" not in employee
    assert admin.put(f"/api/admin/employees/{code}", {"name": "Paid", "department": "D", "designation": "D", "ot_rate": "700.25"}).status_code == 200
    employee = admin.get(f"/api/admin/employees/{code}").json["employee"]
    assert employee["basic_salary"] == 85000.5 and employee["ot_rate"] == 700.25


@pytest.mark.parametrize("bad", ["-1", "abc", "1.234", "NaN", "Infinity", True, "99999999999"])
def test_bad_pay_values_are_refused(admin, bad):
    response = admin.post("/api/admin/employees", {"name": "X", "department": "D", "designation": "D", "basic_salary": bad})
    assert response.status_code == 400 and response.json["field"] == "basic_salary"


def test_pay_is_hidden_without_payroll_permission(make_employee, admin, login_as):
    code = make_employee("Paid", basic_salary=85000, ot_rate=650)
    role = next(r for r in admin.get("/api/admin/roles").json["items"] if r["name"] == "Viewer")
    assert admin.put(f"/api/admin/roles/{role['id']}", {"permissions": ["employees.view", "employees.manage", "backup.manage"]}).status_code == 200
    viewer = login_as("Viewer", "vera")
    for employee in (viewer.get("/api/admin/employees").json["items"][0],
                     viewer.get(f"/api/admin/employees/{code}").json["employee"],
                     viewer.get("/api/admin/export/employees?format=json").json["employees"][0]):
        assert not {"basic_salary", "ot_rate", "basic_salary_cents", "ot_rate_cents"} & set(employee)
    denied = viewer.put(f"/api/admin/employees/{code}", {"name": "Paid", "department": "D", "designation": "D", "basic_salary": 1})
    assert denied.status_code == 403
    assert viewer.put(f"/api/admin/employees/{code}", {"name": "Renamed", "department": "D", "designation": "D"}).status_code == 200
    assert admin.get(f"/api/admin/employees/{code}").json["employee"]["basic_salary"] == 85000


def test_public_endpoints_never_expose_pay(make_employee, anon):
    code = make_employee("Paid", basic_salary=85000, ot_rate=650)
    assert b"85000" not in anon.get("/api/public/employees/search?q=").data
    assert b"85000" not in anon.get(f"/api/public/employees/{code}/status").data


def test_audit_log_does_not_keep_salary_amounts(make_employee, admin):
    code = make_employee("Paid", basic_salary=85000)
    admin.put(f"/api/admin/employees/{code}", {"name": "Paid", "department": "D", "designation": "D", "basic_salary": 91234})
    assert b"91234" not in admin.get("/api/admin/audit").data


# --- Who may add overtime -------------------------------------------------

@pytest.mark.parametrize("role,allowed", [
    ("Manager", True), ("Admin", True), ("Super Admin", True), ("Attendance Manager", False), ("Viewer", False),
])
def test_only_managers_and_above_can_add_overtime(make_employee, login_as, clock, role, allowed):
    code = make_employee("Worker")
    api = login_as(role, "tester")
    expected = 201 if allowed else 403
    assert add_ot(api, [code]).status_code == expected
    assert api.get("/api/admin/overtime").status_code == (200 if allowed else 403)
    assert api.get("/admin/overtime").status_code == (200 if allowed else 302)


def test_not_logged_in_cannot_add_overtime(anon, make_employee):
    assert add_ot(anon, [make_employee("Worker")]).status_code == 401


def test_manager_role_cannot_see_or_change_pay(make_employee, admin, login_as):
    code = make_employee("Worker", ot_rate=500, basic_salary=70000)
    manager = login_as("Manager", "mgr")
    assert add_ot(manager, [code]).status_code == 201
    item = manager.get("/api/admin/overtime").json["items"][0]
    assert "ot_pay" not in item and "rate" not in item and "basic_salary" not in item
    assert manager.get(f"/api/admin/reports/overtime_pay?{RANGE}").status_code == 403
    assert manager.get(f"/api/admin/reports-export?{RANGE}&reports=overtime_pay").status_code == 403
    assert "overtime_pay" not in [r["key"] for r in manager.get("/api/admin/reports").json["items"]]
    assert manager.get(f"/api/admin/reports/added_overtime?{RANGE}").status_code == 200
    book = load_workbook(io.BytesIO(manager.get(f"/api/admin/reports-export?{RANGE}").data))
    assert "Overtime Pay Report" not in book.sheetnames and "Added Overtime Report" in book.sheetnames
    assert manager.get(f"/api/admin/employees/{code}").status_code == 403


# --- Adding overtime ------------------------------------------------------

def test_add_for_selected_and_all(make_employee, admin):
    a, b = make_employee("A"), make_employee("B")
    gone = make_employee("Gone")
    admin.post(f"/api/admin/employees/{gone}/deactivate")
    assert add_ot(admin, [a]).json["count"] == 1
    everyone = admin.post("/api/admin/overtime", {"all": True, "work_date": "2026-10-26", "hours": 1, "reason": "Stocktake"})
    assert everyone.json["count"] == 2
    rows = admin.get("/api/admin/overtime").json["items"]
    assert sorted(r["employee_id"] for r in rows) == sorted([a, a, b])
    assert rows[0]["added_by"] == "admin"
    assert add_ot(admin, [gone]).status_code == 404


def test_add_is_all_or_nothing(make_employee, admin):
    a = make_employee("A")
    assert add_ot(admin, [a, "EM999"]).status_code == 404
    assert admin.get("/api/admin/overtime").json["total"] == 0


@pytest.mark.parametrize("body", [
    {"hours": 0}, {"hours": -1}, {"hours": 25}, {"hours": "abc"}, {"hours": "NaN"}, {"hours": None},
    {"reason": ""}, {"work_date": "2027-01-01"}, {"work_date": "not-a-date"}, {"work_date": ""},
])
def test_overtime_validation(make_employee, admin, body):
    base = {"employee_ids": [make_employee("A")], "work_date": "2026-10-26", "hours": 2, "reason": "Late job"}
    assert admin.post("/api/admin/overtime", {**base, **body}).status_code == 400


def test_daily_limit_and_selection_required(make_employee, admin):
    a = make_employee("A")
    assert add_ot(admin, [a], hours=20).status_code == 201
    assert add_ot(admin, [a], hours=5).json["code"] == "OVERTIME_LIMIT"
    assert admin.post("/api/admin/overtime", {"work_date": "2026-10-26", "hours": 1, "reason": "x"}).status_code == 400


def test_cancel_keeps_the_record_and_stops_counting(app, make_employee, admin):
    a = make_employee("A", ot_rate=600)
    add_ot(admin, [a])
    entry = admin.get("/api/admin/overtime").json["items"][0]
    url = f"/api/admin/overtime/{entry['id']}/cancel"
    assert admin.post(url, {"reason": ""}).status_code == 400
    assert admin.post(url, {"reason": "Entered by mistake"}).status_code == 200
    assert admin.post(url, {"reason": "Again"}).status_code == 404
    assert admin.get("/api/admin/overtime?status=cancelled").json["items"][0]["cancel_reason"] == "Entered by mistake"
    assert admin.get(f"/api/admin/reports/overtime_pay?{RANGE}").json["total"] == 0
    conn = sqlite3.connect(app.config["TIMORA"]["database"])
    with pytest.raises(sqlite3.IntegrityError):
        conn.execute("DELETE FROM overtime_entries")
    conn.close()


def test_pay_uses_the_rate_when_added(make_employee, admin):
    a = make_employee("A", ot_rate=600, basic_salary=70000)
    add_ot(admin, [a], hours=1.5)
    admin.put(f"/api/admin/employees/{a}", {"name": "A", "department": "D", "designation": "D", "ot_rate": 900})
    add_ot(admin, [a], hours=1)
    row = admin.get(f"/api/admin/reports/overtime_pay?{RANGE}").json["items"][0]
    assert row["ot_pay"] == 1800  # 1.5 x 600 + 1 x 900
    assert row["added_seconds"] == int(2.5 * 3600) and row["ot_rate"] == 900 and row["basic_salary"] == 70000


def test_pay_cents_do_not_drift(make_employee, admin):
    a = make_employee("A", ot_rate="333.33")
    for _ in range(3):
        add_ot(admin, [a], hours=1)
    assert admin.get(f"/api/admin/reports/overtime_pay?{RANGE}").json["items"][0]["ot_pay"] == 999.99


# --- Reports show dates beside clock times --------------------------------

def test_clock_in_and_out_have_separate_date_and_time(make_shift, make_employee, clock_action, clock, admin):
    code = make_employee("Night Worker", make_shift("NGT", "22:00", "07:00"), employee_type="Technician")
    clock.set("2026-10-26 22:05:00")
    clock_action(code, "clock_in")
    clock.set("2026-10-27 07:00:00")
    clock_action(code, "clock_out")

    book = load_workbook(io.BytesIO(admin.get(f"/api/admin/reports/daily?{RANGE}&format=xlsx").data))
    sheet = book.active
    header = [c.value for c in sheet[1]]
    for label in ("Clock In Date", "Clock In Time", "Clock Out Date", "Clock Out Time", "Type"):
        assert label in header
    row = dict(zip(header, [c.value for c in sheet[2]]))
    assert str(row["Clock In Date"])[:10] == "2026-10-26" and str(row["Clock In Time"]) == "22:05:00"
    assert str(row["Clock Out Date"])[:10] == "2026-10-27" and str(row["Clock Out Time"]) == "07:00:00"
    assert row["Type"] == "Technician"

    text = admin.get(f"/api/admin/reports/daily?{RANGE}&format=csv").data.decode("utf-8-sig")
    assert "2026-10-26,22:05:00,2026-10-27,07:00:00" in text


def test_clock_in_reports_also_show_clock_out(make_shift, make_employee, clock_action, clock, admin):
    code = make_employee("Late", make_shift("DAY", "09:00", "18:00"))
    clock.set("2026-10-26 09:30:00")
    clock_action(code, "clock_in")
    clock.set("2026-10-26 18:00:00")
    clock_action(code, "clock_out")
    for report in ("late", "early", "overtime", "missing_clockout", "daily"):
        labels = [c["label"] for c in admin.get(f"/api/admin/reports/{report}?{RANGE}").json["columns"]]
        assert {"Clock In Date", "Clock In Time", "Clock Out Date", "Clock Out Time"} <= set(labels), report


def test_report_filters_by_type(make_employee, admin):
    make_employee("Boss", employee_type="Manager")
    make_employee("Tech", employee_type="Technician")
    add_ot(admin, [make_employee("Mop", employee_type="Cleaner")])
    assert admin.get(f"/api/admin/reports/added_overtime?{RANGE}&employee_type=Cleaner").json["total"] == 1
    assert admin.get(f"/api/admin/reports/added_overtime?{RANGE}&employee_type=Manager").json["total"] == 0


# --- Upgrading an existing database ---------------------------------------

def test_older_database_is_upgraded_keeping_its_data(app, make_employee, admin):
    code = make_employee("Existing")
    path = app.config["TIMORA"]["database"]
    conn = sqlite3.connect(path)
    conn.execute("DROP INDEX ix_employees_type")
    for column in ("employee_type", "basic_salary_cents", "ot_rate_cents"):
        conn.execute(f"ALTER TABLE employees DROP COLUMN {column}")
    conn.execute("DROP TABLE overtime_entries")
    conn.execute("UPDATE roles SET permissions = '[\"dashboard.view\"]' WHERE name = 'Admin'")
    conn.execute("DELETE FROM roles WHERE name = 'Manager'")
    conn.execute("PRAGMA user_version = 1")
    conn.commit()
    conn.close()

    init_db(path)

    conn = sqlite3.connect(path)
    assert conn.execute("PRAGMA user_version").fetchone()[0] == SCHEMA_VERSION
    assert conn.execute("SELECT employee_id, employee_type, ot_rate_cents FROM employees").fetchone() == (code, "", 0)
    assert conn.execute("SELECT COUNT(*) FROM overtime_entries").fetchone()[0] == 0
    assert conn.execute("SELECT COUNT(*) FROM roles WHERE name = 'Manager'").fetchone()[0] == 1
    assert "overtime.manage" in conn.execute("SELECT permissions FROM roles WHERE name = 'Admin'").fetchone()[0]
    conn.close()
    assert add_ot(admin, [code]).status_code == 201


def test_upgrade_does_not_undo_later_role_edits(app, admin):
    path = app.config["TIMORA"]["database"]
    role = next(r for r in admin.get("/api/admin/roles").json["items"] if r["name"] == "Admin")
    assert admin.put(f"/api/admin/roles/{role['id']}", {"permissions": ["dashboard.view"]}).status_code == 200
    init_db(path)
    role = next(r for r in admin.get("/api/admin/roles").json["items"] if r["name"] == "Admin")
    assert role["permissions"] == ["dashboard.view"]
