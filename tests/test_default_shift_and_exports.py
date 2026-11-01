"""The 7 AM to 7 AM default shift, and exporting reports to Excel."""
import io
import zipfile
from datetime import date, datetime, timedelta

import pytest
from openpyxl import load_workbook

RANGE = "preset=custom&from=2026-10-26&to=2026-10-26"


@pytest.fixture
def default_shift(admin):
    shifts = admin.get("/api/admin/shifts").json["items"]
    assert len(shifts) == 1
    return shifts[0]


def test_new_install_has_only_the_default_shift(default_shift):
    assert default_shift["code"] == "DEFAULT"
    assert (default_shift["start_time"], default_shift["end_time"]) == ("07:00", "07:00")
    assert default_shift["working_days"] == [0, 1, 2, 3, 4, 5, 6]
    assert default_shift["is_24h"] is True and default_shift["is_overnight"] is False


@pytest.mark.parametrize("now", ["2026-10-26 07:00:00", "2026-10-26 15:00:00", "2026-10-27 06:59:00"])
def test_clock_in_is_open_all_day_and_counts_to_the_7am_day(default_shift, make_employee, clock_action, clock, admin, now):
    code = make_employee(shift_id=default_shift["id"])
    clock.set(now)
    response = clock_action(code, "clock_in").json
    assert response["attendance_status"] == "Present"
    row = admin.get(f"/api/admin/attendance?employee={code}").json["items"][0]
    assert row["attendance_date"] == "2026-10-26" and row["late_seconds"] == 0


def test_short_session_is_not_marked_half_day_or_early(default_shift, make_employee, clock_action, clock, admin):
    code = make_employee(shift_id=default_shift["id"])
    clock.set("2026-10-26 09:00:00")
    clock_action(code, "clock_in")
    clock.set("2026-10-26 17:00:00")
    assert clock_action(code, "clock_out").json["attendance_status"] == "Present"
    row = admin.get(f"/api/admin/attendance?employee={code}").json["items"][0]
    assert row["working_seconds"] == 8 * 3600
    assert row["early_departure_seconds"] == 0 and row["overtime_seconds"] == 0 and row["break_seconds"] == 0


def test_new_day_opens_at_7am(default_shift, make_employee, clock_action, status_of, clock):
    code = make_employee(shift_id=default_shift["id"])
    clock.set("2026-10-26 09:00:00")
    clock_action(code, "clock_in")
    clock.set("2026-10-26 17:00:00")
    clock_action(code, "clock_out")
    clock.set("2026-10-27 06:59:00")
    assert status_of(code)["status"] == "clocked_out"
    clock.set("2026-10-27 07:00:00")
    assert status_of(code)["action"] == "clock_in"


def test_default_shift_can_be_edited_and_saved(default_shift, admin):
    body = {**default_shift, "name": "Renamed"}
    assert admin.put(f"/api/admin/shifts/{default_shift['id']}", body).status_code == 200


@pytest.fixture
def report_data(make_shift, make_employee, clock_action, clock):
    shift = make_shift("DAY", "09:00", "18:00")
    late = make_employee("Alice", shift, department="Ops")
    absent = make_employee("Bob", shift, department="Sales")
    clock.set("2026-10-26 09:30:00")
    clock_action(late, "clock_in")
    clock.set("2026-10-26 20:00:00")
    clock_action(late, "clock_out")
    clock.set("2026-10-27 08:00:00")
    return late, absent, shift


def _workbook(response):
    assert response.status_code == 200, response.data[:200]
    assert "spreadsheetml" in response.content_type
    return load_workbook(io.BytesIO(response.data))


def test_export_all_excel_has_overview_and_every_report(report_data, admin):
    response = admin.get(f"/api/admin/reports-export?{RANGE}")
    assert "timora-all-reports-2026-10-26-to-2026-10-26.xlsx" in response.headers["Content-Disposition"]
    book = _workbook(response)
    assert book.sheetnames == [
        "Overview", "Employee Attendance Summary", "Daily Attendance Report", "Monthly Attendance Report",
        "Late Arrival Report", "Early Departure Report", "Overtime Report", "Missing Clock-Out Report",
        "Absent Employee Report", "Total Working Hours",
    ]
    overview = {row[0]: row[1] for row in book["Overview"].iter_rows(values_only=True) if row[0]}
    assert overview["Period"] == "2026-10-26 to 2026-10-26"
    assert overview["Employees"] == "All"
    assert overview["Late Arrival Report"] == 1 and overview["Absent Employee Report"] == 1


def test_excel_uses_real_dates_and_durations(report_data, admin):
    book = _workbook(admin.get(f"/api/admin/reports-export?{RANGE}&reports=late"))
    sheet = book["Late Arrival Report"]
    header = [c.value for c in sheet[1]]
    row = dict(zip(header, [c.value for c in sheet[2]]))
    assert row["Employee"] == "Alice"
    assert row["Date"].date() == date(2026, 10, 26)
    assert row["Clock In"] == datetime(2026, 10, 26, 9, 30)
    assert sheet.cell(row=2, column=header.index("Late") + 1).number_format == "[h]:mm"
    assert sheet.cell(row=2, column=header.index("Late") + 1).value == timedelta(minutes=30)
    assert sheet.freeze_panes == "A2"


def test_export_all_filters_by_selected_people_and_reports(report_data, admin):
    late, absent, _shift = report_data
    book = _workbook(admin.get(f"/api/admin/reports-export?{RANGE}&employees={absent}&reports=daily,absent"))
    assert book.sheetnames == ["Overview", "Daily Attendance Report", "Absent Employee Report"]
    names = [r[2] for r in book["Daily Attendance Report"].iter_rows(min_row=2, values_only=True)]
    assert names == ["Bob"]
    both = _workbook(admin.get(f"/api/admin/reports-export?{RANGE}&employees={late},{absent}&reports=daily"))
    assert both["Daily Attendance Report"].max_row == 3


def test_export_all_presets_and_validation(report_data, admin, clock):
    clock.set("2026-11-04 12:00:00")
    assert "2026-10-26-to-2026-11-01" in admin.get("/api/admin/reports-export?preset=last_week").headers["Content-Disposition"]
    assert "2026-10-01-to-2026-10-31" in admin.get("/api/admin/reports-export?preset=last_month").headers["Content-Disposition"]
    assert "2026-11-01-to-2026-11-04" in admin.get("/api/admin/reports-export?preset=this_month").headers["Content-Disposition"]
    assert admin.get("/api/admin/reports-export?reports=nope").status_code == 400
    assert admin.get("/api/admin/reports-export?format=doc").status_code == 400
    assert admin.get("/api/admin/reports-export?employees=x';DROP").status_code == 400


def test_export_all_csv_zip_and_json(report_data, admin):
    archive = admin.get(f"/api/admin/reports-export?{RANGE}&format=csv&reports=late,absent")
    assert archive.content_type == "application/zip"
    with zipfile.ZipFile(io.BytesIO(archive.data)) as z:
        assert sorted(z.namelist()) == ["absent-2026-10-26-to-2026-10-26.csv", "late-2026-10-26-to-2026-10-26.csv"]
        assert "Alice" in z.read("late-2026-10-26-to-2026-10-26.csv").decode("utf-8-sig")
    data = admin.get(f"/api/admin/reports-export?{RANGE}&format=json&reports=absent").json
    assert data["reports"]["Absent Employee Report"][0]["Employee"] == "Bob"


def test_single_report_excel(report_data, admin):
    book = _workbook(admin.get(f"/api/admin/reports/late?{RANGE}&format=xlsx"))
    assert book.sheetnames == ["Late Arrival Report"]


def test_excel_never_turns_employee_text_into_formulas(make_shift, make_employee, clock, admin):
    make_employee("=HYPERLINK(\"http://x\")\x07", make_shift("DAY", "09:00", "18:00"))
    clock.set("2026-10-27 08:00:00")
    sheet = _workbook(admin.get(f"/api/admin/reports/absent?{RANGE}&format=xlsx")).active
    cell = sheet.cell(row=2, column=3)
    assert cell.data_type == "s" and cell.value == '=HYPERLINK("http://x")'


def test_export_requires_permission(anon):
    assert anon.get("/api/admin/reports-export").status_code == 401
