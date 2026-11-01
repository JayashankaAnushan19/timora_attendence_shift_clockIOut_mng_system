"""Attendance reports: filters, sorting, pagination, and CSV/JSON/Excel export."""
import csv
import io
import json
import re
import zipfile
from collections import OrderedDict
from datetime import date, timedelta

from flask import Blueprint, Response, request
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import timeutil
from .api import ApiError, check_date, ok, page_args, paginate_list
from .attendance import build_roster, check_range, close_stale_sessions
from .auth import current_admin, require
from .database import get_db
from .overtime import fetch_entries
from .settings import business_tz, get_settings

bp = Blueprint("reports", __name__)

ID = ("employee_id", "Employee ID", "text")
NAME = ("employee_name", "Employee", "text")
TYPE = ("employee_type", "Type", "text")
DEPT = ("department", "Department", "text")
SHIFT = ("shift_name", "Shift", "text")
DATE = ("attendance_date", "Date", "date")
IN = [("clock_in", "Clock In Date", "clock_date"), ("clock_in", "Clock In Time", "clock_time")]
OUT = [("clock_out", "Clock Out Date", "clock_date"), ("clock_out", "Clock Out Time", "clock_time")]
WORKED = ("working_seconds", "Worked", "duration")
LATE = ("late_seconds", "Late", "duration")
EARLY = ("early_departure_seconds", "Early Departure", "duration")
OT = ("overtime_seconds", "Overtime", "duration")
STATUS = ("status", "Status", "status")

COUNTS = [
    ("scheduled_days", "Scheduled", "int"),
    ("days_present", "Present", "int"),
    ("late_count", "Late", "int"),
    ("early_count", "Early Dep.", "int"),
    ("absent_days", "Absent", "int"),
    ("leave_days", "Leave", "int"),
    ("incomplete_count", "Incomplete", "int"),
]
TOTALS = [
    ("working_seconds", "Worked", "duration"),
    ("late_seconds", "Total Late", "duration"),
    ("overtime_seconds", "Overtime", "duration"),
]

XLSX_MIMETYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
EXCEL_FORMATS = {"datetime": "yyyy-mm-dd hh:mm", "date": "yyyy-mm-dd", "clock_date": "yyyy-mm-dd",
                 "clock_time": "hh:mm:ss", "duration": "[h]:mm", "money": "#,##0.00"}
# Only people who may see pay figures can run these.
PAYROLL_REPORTS = {"overtime_pay"}
# Built from added overtime entries rather than the attendance roster.
ENTRY_REPORTS = {"added_overtime", "overtime_pay"}
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="172554")
EMPLOYEE_CODE_RE = re.compile(r"^[A-Za-z]{1,5}\d{3,4}$")


def _empty_summary(row):
    return {
        "employee_id": row["employee_id"], "employee_name": row["employee_name"],
        "employee_type": row["employee_type"], "department": row["department"], "scheduled_days": 0,
        "days_present": 0, "late_count": 0, "early_count": 0, "absent_days": 0, "leave_days": 0,
        "incomplete_count": 0, "working_seconds": 0, "break_seconds": 0, "late_seconds": 0,
        "early_departure_seconds": 0, "overtime_seconds": 0,
    }


def _accumulate(total, row):
    status = row["status"]
    if status == "Absent":
        total["absent_days"] += 1
    elif status == "Leave":
        total["leave_days"] += 1
    if row["clock_in"] or status in ("Absent", "Leave"):
        total["scheduled_days"] += 1
    if row["clock_in"]:
        total["days_present"] += 1
    if status == "Incomplete":
        total["incomplete_count"] += 1
    if row["late_seconds"]:
        total["late_count"] += 1
    if row["early_departure_seconds"]:
        total["early_count"] += 1
    for key in ("working_seconds", "break_seconds", "late_seconds", "early_departure_seconds",
                "overtime_seconds"):
        total[key] += row[key]


def _summary(roster):
    totals = OrderedDict()
    for row in roster:
        _accumulate(totals.setdefault(row["employee_id"], _empty_summary(row)), row)
    return list(totals.values())


def _monthly(roster):
    totals = OrderedDict()
    for row in roster:
        month = row["attendance_date"][:7]
        key = (month, row["employee_id"])
        if key not in totals:
            totals[key] = {"month": month, **_empty_summary(row)}
        _accumulate(totals[key], row)
    return list(totals.values())


def _hours(roster):
    rows = _summary(roster)
    for row in rows:
        days = row["days_present"]
        row["average_seconds"] = row["working_seconds"] // days if days else 0
    return rows


def _daily(roster):
    return roster


def _late(roster):
    return [r for r in roster if r["late_seconds"] > 0]


def _early(roster):
    return [r for r in roster if r["early_departure_seconds"] > 0]


def _overtime(roster):
    return [r for r in roster if r["overtime_seconds"] > 0]


def _missing_clock_out(roster):
    return [r for r in roster if r["clock_in"] and not r["clock_out"] and not r["is_open"]]


def _absent(roster):
    return [r for r in roster if r["status"] == "Absent"]


def _added_overtime(entries):
    return entries


def _overtime_pay(entries):
    totals = OrderedDict()
    for entry in entries:
        total = totals.setdefault(entry["employee_id"], {
            "employee_id": entry["employee_id"], "employee_name": entry["employee_name"],
            "employee_type": entry["employee_type"], "department": entry["department"],
            "basic_salary": entry["basic_salary"], "ot_rate": entry["current_ot_rate"],
            "entries": 0, "added_seconds": 0, "ot_pay": 0,
        })
        total["entries"] += 1
        total["added_seconds"] += entry["added_seconds"]
        total["ot_pay"] = round(total["ot_pay"] + entry["ot_pay"], 2)
    return list(totals.values())


REPORTS = {
    "summary": ("Employee Attendance Summary", [ID, NAME, TYPE, DEPT] + COUNTS + TOTALS, _summary),
    "daily": ("Daily Attendance Report", [DATE, ID, NAME, TYPE, DEPT, SHIFT, *IN, *OUT, WORKED, LATE, EARLY, OT, STATUS], _daily),
    "monthly": ("Monthly Attendance Report", [("month", "Month", "text"), ID, NAME, TYPE, DEPT] + COUNTS + TOTALS, _monthly),
    "late": ("Late Arrival Report", [DATE, ID, NAME, TYPE, DEPT, SHIFT, *IN, *OUT, LATE, STATUS], _late),
    "early": ("Early Departure Report", [DATE, ID, NAME, TYPE, DEPT, SHIFT, *IN, *OUT, EARLY, STATUS], _early),
    "overtime": ("Overtime Report", [DATE, ID, NAME, TYPE, DEPT, SHIFT, *IN, *OUT, OT, WORKED], _overtime),
    "missing_clockout": ("Missing Clock-Out Report", [DATE, ID, NAME, TYPE, DEPT, SHIFT, *IN, *OUT, STATUS], _missing_clock_out),
    "absent": ("Absent Employee Report", [DATE, ID, NAME, TYPE, DEPT, SHIFT], _absent),
    "hours": ("Total Working Hours", [
        ID, NAME, TYPE, DEPT, ("days_present", "Days Worked", "int"), WORKED,
        ("break_seconds", "Breaks", "duration"), ("average_seconds", "Average / Day", "duration"), OT,
    ], _hours),
    "added_overtime": ("Added Overtime Report", [
        ("work_date", "Date", "date"), ID, NAME, TYPE, DEPT, ("added_seconds", "Added OT", "duration"),
        ("reason", "Reason", "text"), ("added_by", "Added By", "text"), ("created_at", "Added On", "datetime"),
    ], _added_overtime),
    "overtime_pay": ("Overtime Pay Report", [
        ID, NAME, TYPE, DEPT, ("basic_salary", "Basic Salary", "money"), ("ot_rate", "Current OT Rate / Hour", "money"),
        ("entries", "OT Entries", "int"), ("added_seconds", "Added OT", "duration"), ("ot_pay", "OT Pay", "money"),
    ], _overtime_pay),
}


def _can_run(report):
    return report not in PAYROLL_REPORTS or "payroll.manage" in current_admin()["permissions"]


def _report_rows(report, roster, entries):
    return REPORTS[report][2](entries if report in ENTRY_REPORTS else roster)


def resolve_range(args, today):
    preset = args.get("preset", "today")
    week_start = today - timedelta(days=today.weekday())
    month_start = today.replace(day=1)
    if preset == "today":
        return today, today
    if preset == "yesterday":
        day = today - timedelta(days=1)
        return day, day
    if preset == "this_week":
        return week_start, today
    if preset == "last_week":
        return week_start - timedelta(days=7), week_start - timedelta(days=1)
    if preset == "this_month":
        return month_start, today
    if preset == "last_month":
        last_day = month_start - timedelta(days=1)
        return last_day.replace(day=1), last_day
    if preset == "custom":
        start = check_date(args.get("from", ""), "From date")
        end = check_date(args.get("to", ""), "To date")
        return date.fromisoformat(start), date.fromisoformat(end)
    raise ApiError("Unknown date range.", code="VALIDATION_ERROR")


def _read_filters(args):
    filters = {
        "employee": args.get("employee", "").strip(),
        "employees": [c.strip() for c in args.get("employees", "").split(",") if c.strip()],
        "department": args.get("department", "").strip(),
        "employee_type": args.get("employee_type", "").strip(),
        "shift": args.get("shift", "").strip(),
    }
    if filters["shift"] and not filters["shift"].isdigit():
        raise ApiError("Invalid shift filter.", code="VALIDATION_ERROR")
    if not all(EMPLOYEE_CODE_RE.match(code) for code in filters["employees"]):
        raise ApiError("Invalid employee filter.", code="VALIDATION_ERROR")
    return filters


def _load_roster(conn, args, settings, now):
    start, end = resolve_range(args, timeutil.local_date(now, business_tz(settings)))
    start, end = start.isoformat(), end.isoformat()
    check_range(start, end)
    filters = _read_filters(args)
    close_stale_sessions(conn, settings, now)
    entries = fetch_entries(conn, start, end, filters)
    return build_roster(conn, start, end, settings, now, filters), entries, start, end, filters


def _sort(rows, columns, args):
    keys = {c[0] for c in columns}
    sort = args.get("sort")
    if sort in keys:
        reverse = args.get("dir") == "desc"
        rows = sorted(rows, key=lambda r: (r.get(sort) is None, str(r.get(sort) or "").lower()
                      if isinstance(r.get(sort), str) else r.get(sort) or 0), reverse=reverse)
    return rows


def _plain_value(value, kind, tz):
    if value is None:
        return ""
    if kind == "datetime":
        return timeutil.format_local_datetime(value, tz)
    if kind == "clock_date":
        return timeutil.format_local_datetime(value, tz)[:10]
    if kind == "clock_time":
        return timeutil.format_local_datetime(value, tz)[11:]
    if kind == "money":
        return f"{value:.2f}"
    if kind == "duration":
        minutes = int(value) // 60
        return f"{minutes // 60}:{minutes % 60:02d}"
    return value


def _safe_csv(value):
    # Stops spreadsheet apps treating employee-entered text as a formula.
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def rows_to_csv(columns, rows, tz):
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([c[1] for c in columns])
    for row in rows:
        writer.writerow([_safe_csv(_plain_value(row.get(c[0]), c[2], tz)) for c in columns])
    return "﻿" + buffer.getvalue()


def _excel_value(value, kind, tz):
    """Real dates and durations, so Excel can sort and add them up."""
    if value is None or value == "":
        return None
    if kind == "datetime":
        return timeutil.parse_iso(value).astimezone(tz).replace(tzinfo=None)
    if kind == "clock_date":
        return timeutil.parse_iso(value).astimezone(tz).date()
    if kind == "clock_time":
        return timeutil.parse_iso(value).astimezone(tz).time()
    if kind == "date":
        return date.fromisoformat(value)
    if kind == "duration":
        return value / 86400
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def _write_cell(sheet, row, column, value, number_format=None):
    cell = sheet.cell(row=row, column=column, value=value)
    if isinstance(value, str):
        cell.data_type = "s"  # employee-entered text must never run as a formula
    if number_format:
        cell.number_format = number_format
    return cell


def _write_report_sheet(sheet, columns, rows, tz):
    widths = [len(c[1]) + 4 for c in columns]
    for index, (_key, label, _kind) in enumerate(columns, 1):
        cell = _write_cell(sheet, 1, index, label)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    for row_number, row in enumerate(rows, 2):
        for index, (key, _label, kind) in enumerate(columns, 1):
            value = _excel_value(row.get(key), kind, tz)
            _write_cell(sheet, row_number, index, value, EXCEL_FORMATS.get(kind))
            if kind == "datetime":
                widths[index - 1] = max(widths[index - 1], 18)
            elif kind in ("date", "clock_date", "clock_time"):
                widths[index - 1] = max(widths[index - 1], 12)
            elif isinstance(value, str):
                widths[index - 1] = max(widths[index - 1], min(len(value) + 2, 40))
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    if rows:
        sheet.auto_filter.ref = sheet.dimensions


def build_workbook(reports, tz, overview=None):
    """reports: [(title, columns, rows)]. overview: [(label, value)] for a first sheet."""
    book = Workbook()
    first = book.active
    if overview is not None:
        first.title = "Overview"
        for row_number, (label, value) in enumerate(overview, 1):
            _write_cell(first, row_number, 1, label).font = Font(bold=True)
            _write_cell(first, row_number, 2, value)
        row_number = len(overview) + 2
        _write_cell(first, row_number, 1, "Report").font = Font(bold=True)
        _write_cell(first, row_number, 2, "Rows").font = Font(bold=True)
        for title, _columns, rows in reports:
            row_number += 1
            _write_cell(first, row_number, 1, title)
            _write_cell(first, row_number, 2, len(rows))
        first.column_dimensions["A"].width = 32
        first.column_dimensions["B"].width = 42
    for position, (title, columns, rows) in enumerate(reports):
        sheet = first if overview is None and position == 0 else book.create_sheet()
        sheet.title = title[:31]
        _write_report_sheet(sheet, columns, rows, tz)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def attachment(content, filename, mimetype):
    response = Response(content, mimetype=mimetype)
    response.headers["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def _plain_rows(columns, rows, tz):
    return [{c[1]: _plain_value(r.get(c[0]), c[2], tz) for c in columns} for r in rows]


@bp.get("/api/admin/reports/<report>")
@require("reports.view")
def run_report(report):
    if report not in REPORTS:
        raise ApiError("Unknown report.", 404, "NOT_FOUND")
    if not _can_run(report):
        raise ApiError("You do not have permission to view pay reports.", 403, "FORBIDDEN")
    conn = get_db()
    settings = get_settings(conn)
    args = request.args
    roster, entries, start, end, _filters = _load_roster(conn, args, settings, timeutil.now_utc())
    title, columns, _builder = REPORTS[report]
    rows = _sort(_report_rows(report, roster, entries), columns, args)
    tz = business_tz(settings)
    export = args.get("format")
    name = f"timora-{report}-{start}-to-{end}"

    if export == "csv":
        return attachment(rows_to_csv(columns, rows, tz), f"{name}.csv", "text/csv; charset=utf-8")
    if export == "xlsx":
        return attachment(build_workbook([(title, columns, rows)], tz), f"{name}.xlsx", XLSX_MIMETYPE)
    if export == "json":
        body = json.dumps({"report": title, "from": start, "to": end, "timezone": settings["timezone"],
                           "rows": _plain_rows(columns, rows, tz)}, indent=2, default=str)
        return attachment(body, f"{name}.json", "application/json")

    column_info = [{"key": c[0], "label": c[1], "kind": c[2]} for c in columns]
    keep = {c[0] for c in columns} | {"id"}
    trim = lambda items: [{k: v for k, v in r.items() if k in keep} for r in items]
    if args.get("all") == "1":  # printable view
        return ok(title=title, columns=column_info, items=trim(rows), total=len(rows), start=start, end=end)
    page, per_page = page_args()
    return ok(title=title, columns=column_info, items=trim(paginate_list(rows, page, per_page)),
              total=len(rows), page=page, per_page=per_page, start=start, end=end)


@bp.get("/api/admin/reports")
@require("reports.view")
def report_list():
    return ok(items=[{"key": key, "title": spec[0]} for key, spec in REPORTS.items() if _can_run(key)])


@bp.get("/api/admin/reports-export")
@require("reports.view")
def export_all():
    """Build several reports for the same people and dates into one download."""
    conn = get_db()
    settings = get_settings(conn)
    args = request.args
    keys = [k for k in args.get("reports", "").split(",") if k] or [k for k in REPORTS if _can_run(k)]
    unknown = [k for k in keys if k not in REPORTS]
    if unknown:
        raise ApiError(f"Unknown report: {unknown[0]}", code="VALIDATION_ERROR")
    if not all(_can_run(k) for k in keys):
        raise ApiError("You do not have permission to view pay reports.", 403, "FORBIDDEN")
    export = args.get("format", "xlsx")
    if export not in ("xlsx", "csv", "json"):
        raise ApiError("Choose Excel, CSV or JSON.", code="VALIDATION_ERROR")

    now = timeutil.now_utc()
    roster, entries, start, end, filters = _load_roster(conn, args, settings, now)
    tz = business_tz(settings)
    reports = [(REPORTS[k][0], REPORTS[k][1], _report_rows(k, roster, entries)) for k in keys]
    name = f"timora-all-reports-{start}-to-{end}"

    if export == "csv":
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
            for key, (_title, columns, rows) in zip(keys, reports):
                archive.writestr(f"{key}-{start}-to-{end}.csv", rows_to_csv(columns, rows, tz).encode("utf-8"))
        return attachment(buffer.getvalue(), f"{name}.zip", "application/zip")
    if export == "json":
        body = {"from": start, "to": end, "timezone": settings["timezone"],
                "reports": {title: _plain_rows(columns, rows, tz) for title, columns, rows in reports}}
        return attachment(json.dumps(body, indent=2, default=str), f"{name}.json", "application/json")

    shift_name = "All"
    if filters["shift"]:
        row = conn.execute("SELECT name FROM shifts WHERE id = ?", (filters["shift"],)).fetchone()
        shift_name = row["name"] if row else "Unknown"
    overview = [
        ("Company", settings["company_name"]),
        ("Period", f"{start} to {end}"),
        ("Timezone", settings["timezone"]),
        ("Employees", ", ".join(filters["employees"]) or filters["employee"] or "All"),
        ("Employee type", filters["employee_type"] or "All"),
        ("Department", filters["department"] or "All"),
        ("Shift", shift_name),
        ("Generated", timeutil.format_local_datetime(timeutil.to_iso(now), tz)),
    ]
    return attachment(build_workbook(reports, tz, overview), f"{name}.xlsx", XLSX_MIMETYPE)
